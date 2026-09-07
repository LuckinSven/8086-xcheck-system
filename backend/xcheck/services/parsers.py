from __future__ import annotations

import csv
import io
import json
import re
import tempfile
import zipfile
from collections.abc import Iterator
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from typing import BinaryIO

import xlrd
from openpyxl import load_workbook

EXCEL_MEMBER_MAX_BYTES = 524_288_000
ACCESS_SOURCE_COLUMNS = ("访问源 IP", "Source IP", "source_ip")
ATTACK_SOURCE_COLUMNS = ("srcAddress", "Source Address", "source_address")
SourceColumns = str | tuple[str, ...]


class ParseError(ValueError):
    pass


@dataclass(frozen=True)
class ExtractedValue:
    raw_value: str
    position: str
    parse_error: str | None = None


def iter_manual_ips(text: str) -> Iterator[ExtractedValue]:
    for index, value in enumerate(re.split(r"[\s,，;；]+", text or ""), start=1):
        value = value.strip()
        if value:
            yield ExtractedValue(value, f"item:{index}")


def _detect_encoding(stream: BinaryIO) -> str:
    position = stream.tell()
    sample = stream.read(131_072)
    stream.seek(position)
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            sample.decode(encoding)
            return encoding
        except UnicodeDecodeError:
            continue
    raise ParseError("文件编码无法识别，支持 UTF-8、UTF-8 BOM 和 GB18030")


def _column_aliases(source_columns: SourceColumns) -> tuple[str, ...]:
    return (source_columns,) if isinstance(source_columns, str) else source_columns


def _source_column_index(header: list | tuple, source_columns: SourceColumns) -> tuple[int, str]:
    normalized = [str(value).strip() if value is not None else "" for value in header]
    for alias in _column_aliases(source_columns):
        if alias in normalized:
            return normalized.index(alias), alias
    raise LookupError


def _missing_column_message(file_type: str, source_columns: SourceColumns) -> str:
    return f"{file_type}缺少固定字段：{' / '.join(_column_aliases(source_columns))}"


def iter_csv_ips(
    stream: BinaryIO,
    source_column: SourceColumns = ACCESS_SOURCE_COLUMNS,
) -> Iterator[ExtractedValue]:
    encoding = _detect_encoding(stream)
    wrapper = io.TextIOWrapper(stream, encoding=encoding, newline="")
    try:
        reader = csv.DictReader(wrapper)
        try:
            _column_index, selected_column = _source_column_index(
                reader.fieldnames or [], source_column
            )
        except LookupError as exc:
            raise ParseError(_missing_column_message("CSV", source_column)) from exc
        for row_number, row in enumerate(reader, start=2):
            value = (row.get(selected_column) or "").strip()
            if value:
                yield ExtractedValue(value, f"row:{row_number}")
            else:
                yield ExtractedValue("", f"row:{row_number}", f"{selected_column}为空")
    except csv.Error as exc:
        raise ParseError(f"CSV结构错误：{exc}") from exc
    finally:
        wrapper.detach()


def iter_jsonl_ips(
    stream: BinaryIO,
    source_column: SourceColumns = ATTACK_SOURCE_COLUMNS,
) -> Iterator[ExtractedValue]:
    encoding = _detect_encoding(stream)
    wrapper = io.TextIOWrapper(stream, encoding=encoding)
    try:
        for line_number, line in enumerate(wrapper, start=1):
            raw_line = line.strip()
            if not raw_line:
                continue
            try:
                item = json.loads(raw_line)
            except json.JSONDecodeError:
                yield ExtractedValue(raw_line[:256], f"line:{line_number}", "JSON格式错误")
                continue
            selected_column = next(
                (
                    alias
                    for alias in _column_aliases(source_column)
                    if isinstance(item, dict) and alias in item
                ),
                None,
            )
            value = item.get(selected_column) if isinstance(item, dict) and selected_column else None
            if not isinstance(value, str) or not value.strip():
                expected = _column_aliases(source_column)[0]
                yield ExtractedValue(raw_line[:256], f"line:{line_number}", f"缺少{expected}")
                continue
            yield ExtractedValue(value.strip(), f"line:{line_number}")
    finally:
        wrapper.detach()


def _source_column(input_type: str) -> tuple[str, ...]:
    if input_type == "attack":
        return ATTACK_SOURCE_COLUMNS
    if input_type in {"access", "csv"}:
        return ACCESS_SOURCE_COLUMNS
    raise ParseError(f"不支持的日志类型：{input_type}")


def _iter_xls_rows(workbook, source_column: SourceColumns) -> Iterator[ExtractedValue]:
    try:
        sheet = workbook.sheet_by_index(0)
        if sheet.nrows == 0:
            raise ParseError("Excel文件没有表头")
        header = [sheet.cell_value(0, index) for index in range(sheet.ncols)]
        try:
            column_index, selected_column = _source_column_index(header, source_column)
        except LookupError as exc:
            raise ParseError(_missing_column_message("Excel", source_column)) from exc
        for row_index in range(1, sheet.nrows):
            value = str(sheet.cell_value(row_index, column_index)).strip()
            if value:
                yield ExtractedValue(value, f"row:{row_index + 1}")
            else:
                yield ExtractedValue("", f"row:{row_index + 1}", f"{selected_column}为空")
    finally:
        workbook.release_resources()


def iter_excel_ips(
    stream: BinaryIO,
    suffix: str,
    source_column: SourceColumns,
    source_path: Path | None = None,
) -> Iterator[ExtractedValue]:
    if suffix == ".xlsx":
        workbook = load_workbook(stream, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            rows = sheet.iter_rows(values_only=True)
            header = next(rows, None)
            try:
                column_index, selected_column = _source_column_index(header or (), source_column)
            except LookupError as exc:
                raise ParseError(_missing_column_message("Excel", source_column)) from exc
            for row_number, row in enumerate(rows, start=2):
                value = row[column_index] if column_index < len(row) else None
                if value is None or not str(value).strip():
                    yield ExtractedValue("", f"row:{row_number}", f"{selected_column}为空")
                else:
                    yield ExtractedValue(str(value).strip(), f"row:{row_number}")
        finally:
            workbook.close()
        return

    try:
        if source_path is not None:
            yield from _iter_xls_rows(
                xlrd.open_workbook(filename=str(source_path), on_demand=True),
                source_column,
            )
            return
        with tempfile.NamedTemporaryFile(suffix=".xls") as temporary:
            while chunk := stream.read(1024 * 1024):
                temporary.write(chunk)
            temporary.flush()
            yield from _iter_xls_rows(
                xlrd.open_workbook(filename=temporary.name, on_demand=True),
                source_column,
            )
    except xlrd.XLRDError as exc:
        raise ParseError(f"XLS文件损坏：{exc}") from exc


def iter_zip_ips(
    path: Path,
    max_members: int,
    max_uncompressed_bytes: int,
    input_type: str = "attack",
) -> Iterator[ExtractedValue]:
    try:
        archive = zipfile.ZipFile(path)
    except zipfile.BadZipFile as exc:
        raise ParseError("ZIP文件损坏") from exc
    with archive:
        files = [member for member in archive.infolist() if not member.is_dir()]
        if len(files) > max_members:
            raise ParseError(f"ZIP成员数量超过限制：{max_members}")
        if sum(member.file_size for member in files) > max_uncompressed_bytes:
            raise ParseError("ZIP解压后累计大小超过限制")
        supported = 0
        for member in files:
            pure_path = PurePosixPath(member.filename)
            if pure_path.is_absolute() or ".." in pure_path.parts:
                raise ParseError("ZIP包含非法路径")
            suffix = pure_path.suffix.lower()
            supported_suffixes = {".csv", ".xls", ".xlsx", ".log", ".jsonl"}
            if suffix not in supported_suffixes:
                continue
            if suffix in {".xls", ".xlsx"} and member.file_size > EXCEL_MEMBER_MAX_BYTES:
                raise ParseError("ZIP内Excel成员大小超过500 MB限制")
            supported += 1
            with archive.open(member) as stream:
                column = _source_column(input_type)
                if suffix == ".csv":
                    iterator = iter_csv_ips(stream, column)
                elif suffix in {".xls", ".xlsx"}:
                    iterator = iter_excel_ips(stream, suffix, column)
                else:
                    iterator = iter_jsonl_ips(stream, column)
                for item in iterator:
                    yield ExtractedValue(
                        item.raw_value, f"{member.filename}:{item.position}", item.parse_error
                    )
        if supported == 0:
            raise ParseError("ZIP中没有受支持的 CSV、XLS、XLSX、LOG 或 JSONL 文件")


def iter_file_ips(
    path: Path,
    input_type: str,
    max_members: int,
    max_uncompressed_bytes: int,
) -> Iterator[ExtractedValue]:
    suffix = path.suffix.lower()
    if suffix == ".zip":
        yield from iter_zip_ips(path, max_members, max_uncompressed_bytes, input_type)
        return

    column = _source_column(input_type)
    with path.open("rb") as stream:
        if suffix == ".csv":
            yield from iter_csv_ips(stream, column)
        elif suffix in {".xls", ".xlsx"}:
            yield from iter_excel_ips(stream, suffix, column, path)
        elif suffix in {".log", ".jsonl"}:
            yield from iter_jsonl_ips(stream, column)
        else:
            raise ParseError("无法识别文件格式，支持 CSV、XLS、XLSX、ZIP、LOG 和 JSONL")
