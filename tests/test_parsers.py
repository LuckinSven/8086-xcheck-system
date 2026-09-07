import io
import json
import zipfile

import pytest
import xlwt
from openpyxl import Workbook
from xcheck.services import parsers
from xcheck.services.parsers import (
    ParseError,
    iter_csv_ips,
    iter_excel_ips,
    iter_file_ips,
    iter_jsonl_ips,
    iter_manual_ips,
    iter_zip_ips,
)


def test_manual_parser_accepts_common_separators():
    values = list(iter_manual_ips("8.8.8.8,1.1.1.1；2001:4860:4860::8888\n错误地址"))
    assert [item.raw_value for item in values] == [
        "8.8.8.8",
        "1.1.1.1",
        "2001:4860:4860::8888",
        "错误地址",
    ]


def test_csv_parser_streams_fixed_chinese_column():
    content = "访问源 IP,时间\n8.8.8.8,2026-08-13\n1.1.1.1,2026-08-13\n".encode()
    values = list(iter_csv_ips(io.BytesIO(content)))
    assert [(item.raw_value, item.position) for item in values] == [
        ("8.8.8.8", "row:2"),
        ("1.1.1.1", "row:3"),
    ]


def test_csv_parser_rejects_missing_source_column():
    with pytest.raises(ParseError, match="访问源 IP"):
        list(iter_csv_ips(io.BytesIO("客户端,时间\n8.8.8.8,t\n".encode())))


def test_jsonl_parser_records_bad_lines_and_continues():
    content = b'{"srcAddress":"8.8.8.8"}\nnot-json\n{"srcAddress":"1.1.1.1"}\n'
    values = list(iter_jsonl_ips(io.BytesIO(content)))
    assert [item.raw_value for item in values] == ["8.8.8.8", "not-json", "1.1.1.1"]
    assert values[1].parse_error == "JSON格式错误"


def test_zip_parser_reads_supported_member_and_rejects_traversal(tmp_path):
    safe = tmp_path / "safe.zip"
    with zipfile.ZipFile(safe, "w") as archive:
        archive.writestr("attack.log", json.dumps({"srcAddress": "8.8.8.8"}) + "\n")
    assert [item.raw_value for item in iter_zip_ips(safe, 20, 1024)] == ["8.8.8.8"]

    unsafe = tmp_path / "unsafe.zip"
    with zipfile.ZipFile(unsafe, "w") as archive:
        archive.writestr("../attack.log", "{}")
    with pytest.raises(ParseError, match="非法路径"):
        list(iter_zip_ips(unsafe, 20, 1024))


def test_business_type_selects_column_in_csv(tmp_path):
    attack = tmp_path / "attack.csv"
    attack.write_text("srcAddress,event\n8.8.8.8,scan\n", encoding="utf-8")
    access = tmp_path / "access.csv"
    access.write_text("访问源 IP,时间\n1.1.1.1,now\n", encoding="utf-8")

    assert [item.raw_value for item in iter_file_ips(attack, "attack", 20, 1024)] == ["8.8.8.8"]
    assert [item.raw_value for item in iter_file_ips(access, "access", 20, 1024)] == ["1.1.1.1"]


@pytest.mark.parametrize("header", ["访问源 IP", "Source IP", "source_ip"])
def test_access_csv_accepts_bilingual_source_column_aliases(tmp_path, header):
    path = tmp_path / f"access-{header.replace(' ', '-')}.csv"
    path.write_text(f"{header},time\n1.1.1.1,now\n", encoding="utf-8")

    assert [item.raw_value for item in iter_file_ips(path, "access", 20, 1024)] == ["1.1.1.1"]


@pytest.mark.parametrize("header", ["srcAddress", "Source Address", "source_address"])
def test_attack_jsonl_accepts_bilingual_source_column_aliases(tmp_path, header):
    path = tmp_path / f"attack-{header.replace(' ', '-')}.jsonl"
    path.write_text(json.dumps({header: "8.8.8.8"}) + "\n", encoding="utf-8")

    assert [item.raw_value for item in iter_file_ips(path, "attack", 20, 1024)] == ["8.8.8.8"]


def test_excel_accepts_english_source_column_alias(tmp_path):
    path = tmp_path / "access-english.xlsx"
    workbook = Workbook()
    workbook.active.append(["source_ip", "time"])
    workbook.active.append(["9.9.9.9", "now"])
    workbook.save(path)

    assert [item.raw_value for item in iter_file_ips(path, "access", 20, 1024)] == ["9.9.9.9"]


def test_xlsx_and_xls_use_business_column(tmp_path):
    xlsx_path = tmp_path / "attack.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["srcAddress", "event"])
    sheet.append(["8.8.4.4", "scan"])
    workbook.save(xlsx_path)

    xls_path = tmp_path / "access.xls"
    legacy = xlwt.Workbook()
    legacy_sheet = legacy.add_sheet("Sheet1")
    legacy_sheet.write(0, 0, "访问源 IP")
    legacy_sheet.write(0, 1, "时间")
    legacy_sheet.write(1, 0, "1.0.0.1")
    legacy_sheet.write(1, 1, "now")
    legacy.save(str(xls_path))

    assert [item.raw_value for item in iter_file_ips(xlsx_path, "attack", 20, 1024)] == [
        "8.8.4.4"
    ]
    assert [item.raw_value for item in iter_file_ips(xls_path, "access", 20, 1024)] == [
        "1.0.0.1"
    ]


def test_xls_stream_never_requests_an_unbounded_memory_read():
    content = io.BytesIO()
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("Sheet1")
    sheet.write(0, 0, "访问源 IP")
    sheet.write(1, 0, "8.8.8.8")
    workbook.save(content)

    class BoundedReadStream(io.BytesIO):
        def read(self, size=-1):
            assert size >= 0, "XLS parser requested an unbounded read"
            return super().read(size)

    values = list(iter_excel_ips(BoundedReadStream(content.getvalue()), ".xls", "访问源 IP"))
    assert [item.raw_value for item in values] == ["8.8.8.8"]


def test_zip_uses_task_business_type_for_mixed_supported_members(tmp_path):
    archive_path = tmp_path / "access.zip"
    with zipfile.ZipFile(archive_path, "w") as archive:
        archive.writestr("access.csv", "访问源 IP,时间\n9.9.9.9,now\n")
        archive.writestr("ignored.jsonl", json.dumps({"srcAddress": "8.8.8.8"}) + "\n")

    values = list(iter_file_ips(archive_path, "access", 20, 4096))
    assert values[0].raw_value == "9.9.9.9"
    assert values[1].parse_error == "缺少访问源 IP"


def test_access_zip_accepts_jsonl_and_log_members(tmp_path):
    archive_path = tmp_path / "access-json.zip"
    with zipfile.ZipFile(archive_path, "w") as archive:
        archive.writestr("access.jsonl", json.dumps({"访问源 IP": "8.8.8.8"}) + "\n")
        archive.writestr("access.log", json.dumps({"访问源 IP": "1.1.1.1"}) + "\n")

    values = list(iter_file_ips(archive_path, "access", 20, 4096))
    assert [item.raw_value for item in values] == ["8.8.8.8", "1.1.1.1"]


def test_zip_rejects_an_individually_oversized_excel_member(tmp_path, monkeypatch):
    source = tmp_path / "access.xlsx"
    workbook = Workbook()
    workbook.active.append(["访问源 IP"])
    workbook.active.append(["8.8.8.8"])
    workbook.save(source)
    archive_path = tmp_path / "large-excel.zip"
    with zipfile.ZipFile(archive_path, "w") as archive:
        archive.write(source, "access.xlsx")
    monkeypatch.setattr(parsers, "EXCEL_MEMBER_MAX_BYTES", 100, raising=False)

    with pytest.raises(ParseError, match="Excel成员大小"):
        list(iter_file_ips(archive_path, "access", 20, 1_048_576))
